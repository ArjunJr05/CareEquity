import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0]
    if level == 1:
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(29, 78, 216) # #1d4ed8 Primary Blue
        run.font.bold = True
    elif level == 2:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(30, 58, 138) # #1e3a8a Dark Blue
        run.font.bold = True
    elif level == 3:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(71, 85, 105) # Slate
        run.font.bold = True
    return h

# ---------------------------------------------------------
# DOCUMENT 1: BACKEND ARCHITECTURE & INTEGRATION DOCUMENT
# ---------------------------------------------------------
doc_backend = docx.Document()

# Page Setup
for section in doc_backend.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Title Block
title_p = doc_backend.add_paragraph()
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after = Pt(2)
title_run = title_p.add_run("CareEquity Backend Architecture & Microservices Specification")
title_run.font.size = Pt(24)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(30, 58, 138)

sub_p = doc_backend.add_paragraph()
sub_p.paragraph_format.space_after = Pt(18)
sub_run = sub_p.add_run("Comprehensive Technical Documentation of Microservices, Data Pipeline, & AI Engines")
sub_run.font.size = Pt(12)
sub_run.font.italic = True
sub_run.font.color.rgb = RGBColor(100, 116, 139)

add_styled_heading(doc_backend, "1. Executive Summary & Architecture Overview", level=1)
p = doc_backend.add_paragraph(
    "CareEquity is a multi-tier, AI-driven healthcare platform designed to analyze Social Determinants of Health (SDoH), "
    "predict individual disease risks, synthesize clinical care plans, and connect care coordinators with verified non-profit resources. "
    "The backend architecture is built as a microservice cluster comprising 6 core microservices operating synchronously via REST APIs."
)
p.paragraph_format.space_after = Pt(8)

# Table of Microservices
table = doc_backend.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = table.rows[0].cells
headers = ["Service Name", "Container / Port", "Technology Stack", "Core Responsibility"]
for idx, text in enumerate(headers):
    hdr_cells[idx].text = text
    set_cell_background(hdr_cells[idx], "1D4ED8")
    p = hdr_cells[idx].paragraphs[0]
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
    p.runs[0].font.size = Pt(10)

services_data = [
    ("Main System Backend", "careequity-main-backend : 8000", "FastAPI / PostgreSQL", "Patient CRUD, user authentication, subscription management, payment processing (Razorpay), history tracking."),
    ("OCR Data Extraction Engine", "careequity-ocr-backend : 8001", "FastAPI / PyPDF2 / pdfplumber", "Parses uploaded clinical PDFs/documents, extracts structured JSON demographics and lab results."),
    ("RAG & Chatbot Engine", "careequity-rag-backend : 8002", "FastAPI / Groq LLM / ChromaDB", "RAG vector retrieval over census SDoH data and PubMed clinical literature for Consult AI Chatbot."),
    ("4-Agent Research Assistant", "careequity-agent-backend : 8003", "FastAPI / Multi-Agent Framework", "Orchestrates 4 specialized agents (Clinical, SDoH, Equity, Policy) to synthesize integrated care plans."),
    ("Knowledge Graph (KG) Engine", "careequity-kg-backend : 8004", "FastAPI / Neo4j Graph DB", "Graph relationships between counties, risk drivers, SVI scores, and intervention strategies."),
    ("NGO Connect Engine", "careequity-ngo-backend : 8005", "FastAPI / Haversine Geo / SMTP", "Geospatial matching of top 3 nearest verified non-profits per intervention domain & SMTP email dispatch.")
]

for name, port, tech, desc in services_data:
    row_cells = table.add_row().cells
    row_cells[0].text = name
    row_cells[1].text = port
    row_cells[2].text = tech
    row_cells[3].text = desc
    for cell in row_cells:
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        for p in cell.paragraphs:
            p.runs[0].font.size = Pt(9.5)

add_styled_heading(doc_backend, "2. Key Backend Microservice Details", level=1)

add_styled_heading(doc_backend, "2.1 ML Disease Risk Prediction Model (ml_pipelineV2.pkl)", level=2)
doc_backend.add_paragraph(
    "The Machine Learning subsystem utilizes a Scikit-Learn multi-label classification pipeline trained on 75,000 synthetic clinical records "
    "merged with CDC SVI (Social Vulnerability Index) county census data. It evaluates 19 medical features and 11 SDoH environmental features "
    "to compute individual probabilities for Diabetes, Hypertension, Heart Disease, and Asthma."
)

add_styled_heading(doc_backend, "2.2 4-Agent Synthesis Engine", level=2)
doc_backend.add_paragraph(
    "The research assistant backend orchestrates four autonomous AI agents in sequence:\n"
    "• Agent 1 (Clinical Specialist): Analyzes patient vitals, HbA1c, BP, and disease probabilities.\n"
    "• Agent 2 (SDoH Strategist): Evaluates county-level poverty, food insecurity, and transportation barriers.\n"
    "• Agent 3 (Health Equity Advocate): Formulates culturally tailored outreach strategies.\n"
    "• Agent 4 (Policy & Interventions Director): Synthesizes findings into a unified, downloadable clinical care report."
)

add_styled_heading(doc_backend, "2.3 NGO Connect & Email Dispatch Subsystem", level=2)
doc_backend.add_paragraph(
    "Integrated on port 8005, the NGO Connect engine executes Haversine distance calculations against a database of 1,275 verified non-profits. "
    "When a care coordinator selects an intervention (e.g. Housing, Food, Healthcare), the service calculates proximity from patient coordinates "
    "and supports direct SMTP email dispatch to facilitate resource referral."
)

doc_backend.save("e:/CareEquity/CareEquity_Backend_Documentation.docx")
print("Saved CareEquity_Backend_Documentation.docx")

# ---------------------------------------------------------
# DOCUMENT 2: DEPLOYMENT & DEVOPS SPECIFICATION DOCUMENT
# ---------------------------------------------------------
doc_deploy = docx.Document()

for section in doc_deploy.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Title Block
title_p = doc_deploy.add_paragraph()
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after = Pt(2)
title_run = title_p.add_run("CareEquity AWS EC2 Deployment & DevOps Guide")
title_run.font.size = Pt(24)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(16, 185, 129) # Emerald Green

sub_p = doc_deploy.add_paragraph()
sub_p.paragraph_format.space_after = Pt(18)
sub_run = sub_p.add_run("Complete Reference for Production Deployment, GitHub Actions CI/CD, & Docker Stack")
sub_run.font.size = Pt(12)
sub_run.font.italic = True
sub_run.font.color.rgb = RGBColor(100, 116, 139)

add_styled_heading(doc_deploy, "1. Infrastructure Overview", level=1)
doc_deploy.add_paragraph(
    "CareEquity is hosted on an AWS EC2 Ubuntu instance (IP: 18.60.232.212 / Region: ap-south-2). "
    "The application relies on Docker Compose orchestration to manage 7 distinct containers, connected to AWS RDS PostgreSQL "
    "and external vector/graph databases."
)

# Table of Deployment Config
table_d = doc_deploy.add_table(rows=1, cols=3)
table_d.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = table_d.rows[0].cells
headers = ["Component", "Configuration", "Description / Port"]
for idx, text in enumerate(headers):
    hdr_cells[idx].text = text
    set_cell_background(hdr_cells[idx], "059669")
    p = hdr_cells[idx].paragraphs[0]
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
    p.runs[0].font.size = Pt(10)

deploy_specs = [
    ("AWS EC2 Instance", "t3.small / Ubuntu 22.04 LTS", "Public IP: 18.60.232.212"),
    ("Database Layer", "AWS RDS PostgreSQL", "Port 5432 (caredb.cxcssmmegdh1.ap-south-2.rds.amazonaws.com)"),
    ("Security Group Rules", "Inbound Rules sg-01a08f2be3f5ef2aa", "HTTP (80), HTTPS (443), SSH (22), Frontend (5173), Backend API Range (8000-8005)"),
    ("CI/CD Automation", "GitHub Actions Self-Hosted Runner", "Triggered on git push origin main"),
    ("Container Runtime", "Docker & Docker Compose", "Multi-stage builds, isolated networks, automated volume mounting")
]

for comp, config, desc in deploy_specs:
    row_cells = table_d.add_row().cells
    row_cells[0].text = comp
    row_cells[1].text = config
    row_cells[2].text = desc
    for cell in row_cells:
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        for p in cell.paragraphs:
            p.runs[0].font.size = Pt(9.5)

add_styled_heading(doc_deploy, "2. GitHub Actions CI/CD Pipeline (.github/workflows/main.yml)", level=1)
doc_deploy.add_paragraph(
    "Automated continuous integration and deployment is executed on every push to the main branch via a self-hosted runner on the EC2 server:\n\n"
    "1. Environment Injection: Dynamically writes secrets (GROQ_API_KEY, NVIDIA_API_KEY, NEO4J_URI, RAZORPAY_KEY_ID, DATABASE_URL) to .env files.\n"
    "2. EBS Storage Auto-Expansion: Automatically resizes EC2 storage partitions to utilize full 30GB EBS disk space.\n"
    "3. Automated System Pruning: Prunes dangling Docker images, build caches, and system temp folders prior to compilation.\n"
    "4. Memory OOM Protection: Dynamically configures a 2GB swapfile (/swapfile) to prevent Out-Of-Memory (Exit Code 137) build failures.\n"
    "5. Sequential Container Compilation: Executes 'docker compose build' sequentially to maintain low RAM pressure during image generation.\n"
    "6. Service Orchestration: Launches containers via 'docker compose up -d --remove-orphans' and validates health status."
)

add_styled_heading(doc_deploy, "3. Security Group & Port Reference", level=1)
doc_deploy.add_paragraph(
    "The AWS Security Group sg-01a08f2be3f5ef2aa enforces inbound access control:\n"
    "• 80 / 443 (HTTP/HTTPS): Public web traffic.\n"
    "• 22 (SSH): Secure remote administration.\n"
    "• 8000 - 8005 (Custom TCP Range): Public REST API endpoints for Main, OCR, RAG, Agent, KG, and NGO Connect services."
)

doc_deploy.save("e:/CareEquity/CareEquity_Deployment_Documentation.docx")
print("Saved CareEquity_Deployment_Documentation.docx")
